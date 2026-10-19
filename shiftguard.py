#!/usr/bin/env python3
"""ShiftGuard v2: configurable SOC roster and real per-shift coverage validation.

This module uses synthetic sample staff. Never publish real staff details in a public repo.
"""
from __future__ import annotations
import argparse
import calendar
import csv
import json
from collections import Counter, defaultdict
from datetime import date, datetime, timedelta
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

SHIFTS = ("Morning", "Afternoon", "Night")
COLORS = {"Morning": "DCEBFF", "Afternoon": "DEF6E8", "Night": "EEE4FE",
          "Week Off": "F1F3F7", "General": "FFF2D7", "Not Active": "EEEEEE"}
HEAD = "173454"
DEFAULT_REQUIREMENTS = {"SOC L1": 1, "SOC L2": 1}
SAMPLE = [
 {"name": f"SOC L1 Analyst {n}", "role": "SOC L1", "offset": n - 1, "pattern": "rotating"}
 for n in range(1, 6)
] + [
 {"name": f"SOC L2 Analyst {n}", "role": "SOC L2", "offset": n - 1, "pattern": "rotating"}
 for n in range(1, 6)
] + [{"name": "SOC Lead (Sample)", "role": "Team Lead", "pattern": "weekday"}]

def parse_day(value: str) -> date:
    return date.fromisoformat(value)

def month_days(year: int, month: int):
    return [date(year, month, d) for d in range(1, calendar.monthrange(year, month)[1] + 1)]

def validate_team(team):
    if not isinstance(team, list) or not team:
        raise ValueError("Team must be a non-empty array")
    seen = set()
    for item in team:
        if not isinstance(item, dict):
            raise ValueError("Every team member must be an object")
        for field in ("name", "role", "pattern"):
            if not isinstance(item.get(field), str) or not item[field].strip():
                raise ValueError(f"Team member missing valid {field}")
        if item["name"] in seen:
            raise ValueError(f"Duplicate name: {item['name']}")
        seen.add(item["name"])
        if item["pattern"] not in ("rotating", "weekday"):
            raise ValueError("Pattern must be rotating or weekday")
        if not isinstance(item.get("offset", 0), int) or item.get("offset", 0) < 0:
            raise ValueError("offset must be a non-negative integer")
        for key in ("start", "end"):
            if key in item:
                parse_day(item[key])
        if item.get("start") and item.get("end") and item["start"] > item["end"]:
            raise ValueError("start must be before end")

def schedule(team, days, cycle_length=5):
    """Fixed 3-on/2-off rotation over five days, or 3-on/1-off over four days.

    Position anchored to 2026-01-01 to prevent month-boundary reset.
    """
    validate_team(team)
    if cycle_length not in (4, 5):
        raise ValueError("cycle_length must be 4 or 5")
    base = date(2026, 1, 1)
    result = {}
    for member in team:
        shifts = {}
        for day in days:
            if (member.get("start") and day < parse_day(member["start"])) or (
                member.get("end") and day > parse_day(member["end"])
            ):
                assignment = "Not Active"
            elif member["pattern"] == "weekday":
                assignment = "General" if day.weekday() < 5 else "Week Off"
            else:
                idx = ((day - base).days + member.get("offset", 0)) % cycle_length
                assignment = SHIFTS[idx] if idx < 3 else "Week Off"
            shifts[day.isoformat()] = assignment
        result[member["name"]] = shifts
    return result

def coverage(team, days, assigned, requirements):
    """Check scheduled analysts in EACH named shift and role, not headcount."""
    output = []
    for day in days:
        for role, minimum in requirements.items():
            for shift in SHIFTS:
                people = [m["name"] for m in team
                          if m["role"] == role and assigned[m["name"]][day.isoformat()] == shift]
                if len(people) < minimum:
                    output.append({"date": day.isoformat(), "role": role, "shift": shift,
                                   "scheduled": len(people), "required": minimum,
                                   "gap": minimum - len(people)})
    return output

def build_workbook(team, days, assigned, gaps, output):
    wb = Workbook()
    ws = wb.active
    ws.title = "SOC Roster"
    ws.append(["Team member", "Role", "Pattern"] + [d.isoformat() for d in days])
    ws.freeze_panes = "D2"
    ws.column_dimensions["A"].width = 30
    ws.column_dimensions["B"].width = 18
    ws.column_dimensions["C"].width = 16
    ws.auto_filter.ref = f"A1:{get_column_letter(len(days)+3)}{len(team)+1}"
    for idx, day in enumerate(days, 4):
        ws.column_dimensions[get_column_letter(idx)].width = 15
    for member in team:
        name = member["name"]
        ws.append([name, member["role"], member["pattern"]]
                  + [assigned[name][d.isoformat()] for d in days])
        row = ws.max_row
        for col in range(4, ws.max_column + 1):
            cell = ws.cell(row, col)
            cell.fill = PatternFill("solid", fgColor=COLORS[cell.value])
            cell.alignment = Alignment(horizontal="center")
    for cell in ws[1]:
        cell.fill = PatternFill("solid", fgColor=HEAD)
        cell.font = Font(bold=True, color="FFFFFF")
    gv = wb.create_sheet("Coverage Gaps")
    gv.append(["Date", "Role", "Shift", "Scheduled", "Required", "Gap"])
    for g in gaps:
        gv.append([g["date"], g["role"], g["shift"], g["scheduled"], g["required"], g["gap"]])
    for cell in gv[1]:
        cell.fill = PatternFill("solid", fgColor=HEAD)
        cell.font = Font(bold=True, color="FFFFFF")
    gv.freeze_panes = "A2"
    for letter, size in zip("ABCDEF", [17, 20, 19, 17, 16, 12]):
        gv.column_dimensions[letter].width = size
    stats = wb.create_sheet("Shift Statistics")
    stats.append(["Name", "Role", "Morning", "Afternoon", "Night", "General", "Week Off", "Not Active"])
    for member in team:
        c = Counter(assigned[member["name"]].values())
        stats.append([member["name"], member["role"]] +
                     [c[k] for k in ("Morning", "Afternoon", "Night", "General", "Week Off", "Not Active")])
    for cell in stats[1]:
        cell.fill = PatternFill("solid", fgColor=HEAD)
        cell.font = Font(bold=True, color="FFFFFF")
    stats.freeze_panes = "C2"
    for col in range(1, 9):
        stats.column_dimensions[get_column_letter(col)].width = 29 if col == 1 else 19
    meta = wb.create_sheet("Read Me")
    meta.append(["ShiftGuard v2 — sample roster; verify coverage against operational requirements"])
    meta.append(["Coverage gaps are based on actual daily shift assignments, not active employee totals."])
    meta.append(["This file does not prove compliance with labor regulations."])
    meta.append(["Times, rest breaks, leave, skill requirements and holidays need separate review."])
    meta.column_dimensions["A"].width = 100
    Path(output).parent.mkdir(parents=True, exist_ok=True)
    wb.save(output)

def main(argv=None):
    parser = argparse.ArgumentParser(description="Create a synthetic or supplied SOC roster and validate each shift.")
    parser.add_argument("--year", type=int, default=2026)
    parser.add_argument("--month", type=int, default=10)
    parser.add_argument("--team", type=Path, help="JSON array of staff records; never commit sensitive staff data")
    parser.add_argument("--cycle", type=int, choices=[4,5], default=5)
    parser.add_argument("--min-l1", type=int, default=1)
    parser.add_argument("--min-l2", type=int, default=1)
    parser.add_argument("--output", type=Path, default=Path("SOC_Roster_Validated.xlsx"))
    parser.add_argument("--fail-on-gaps", action="store_true", help="Exit 2 if minimum shift coverage is not met")
    args = parser.parse_args(argv)
    if args.min_l1 < 0 or args.min_l2 < 0:
        parser.error("Coverage minimum cannot be negative")
    team = json.loads(args.team.read_text(encoding="utf-8")) if args.team else SAMPLE
    days = month_days(args.year, args.month)
    assigned = schedule(team, days, args.cycle)
    req = {"SOC L1": args.min_l1, "SOC L2": args.min_l2}
    gaps = coverage(team, days, assigned, req)
    build_workbook(team, days, assigned, gaps, args.output)
    print(f"Roster: {args.output}; team: {len(team)}; days: {len(days)}; uncovered role/shift slots: {len(gaps)}")
    if gaps:
        print("WARNING: shift coverage gaps exist; see Coverage Gaps worksheet")
    if gaps and args.fail_on_gaps:
        return 2
    return 0

if __name__ == "__main__":
    raise SystemExit(main())
