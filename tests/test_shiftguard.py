import unittest
from datetime import date
from pathlib import Path
import tempfile
from openpyxl import load_workbook
from shiftguard import SAMPLE, build_workbook, coverage, month_days, schedule, validate_team

class ShiftGuardTests(unittest.TestCase):
    def test_exact_month(self):
        self.assertEqual(len(month_days(2028, 2)), 29)
        self.assertEqual(len(month_days(2027, 2)), 28)

    def test_gap_detection_checks_real_shifts(self):
        team = [{"name": "L1", "role": "SOC L1", "pattern": "weekday"}]
        days = [date(2026, 10, 19)]
        assignments = schedule(team, days)
        gaps = coverage(team, days, assignments, {"SOC L1": 1})
        self.assertEqual(len(gaps), 3)
        self.assertEqual({g["shift"] for g in gaps}, {"Morning", "Afternoon", "Night"})

    def test_rotation_independent_of_month_boundary(self):
        team = [{"name": "A", "role": "SOC L1", "pattern": "rotating", "offset": 0}]
        d = [date(2026, 1, 31), date(2026, 2, 1)]
        self.assertEqual(schedule(team,d)["A"], {k.isoformat(): v for k,v in zip(d,["Morning","Afternoon"])})

    def test_duplicate_names_rejected(self):
        with self.assertRaises(ValueError):
            validate_team([{"name":"X","role":"SOC L1","pattern":"rotating"}]*2)

    def test_workbook_has_real_gap_sheet(self):
        days = month_days(2026, 10)
        assignments = schedule(SAMPLE, days)
        gaps = coverage(SAMPLE, days, assignments, {"SOC L1":1,"SOC L2":1})
        with tempfile.TemporaryDirectory() as d:
            output = Path(d)/"roster.xlsx"
            build_workbook(SAMPLE,days,assignments,gaps,output)
            wb=load_workbook(output)
            self.assertEqual(wb.sheetnames,["SOC Roster","Coverage Gaps","Shift Statistics","Read Me"])
            self.assertEqual(wb["Coverage Gaps"].max_row,len(gaps)+1)
if __name__=="__main__":
    unittest.main()
